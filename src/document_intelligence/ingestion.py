"""Document loading and structure-aware chunking."""

from __future__ import annotations

import hashlib
import re
from bisect import bisect_right
from collections.abc import Iterable
from pathlib import Path

from .language import detect_languages
from .models import Chunk, Document

SUPPORTED_SUFFIXES = {".txt", ".md", ".pdf", ".docx"}


def _stable_id(*parts: str) -> str:
    digest = hashlib.sha256("\x1f".join(parts).encode("utf-8")).hexdigest()
    return digest[:16]


def document_from_text(
    name: str,
    text: str,
    *,
    language: str | None = None,
    metadata: dict | None = None,
) -> Document:
    clean_text = text.strip()
    detected_languages = detect_languages(clean_text)
    if language and language != "mixed":
        languages = (language,) if language != "unknown" else detected_languages
        language_label = language
    else:
        languages = detected_languages
        language_label = (
            "mixed" if len(languages) > 1 else languages[0] if languages else "unknown"
        )
    document_metadata = dict(metadata or {})
    document_metadata["languages"] = list(languages)
    return Document(
        id=_stable_id(name, clean_text),
        name=name,
        text=clean_text,
        language=language_label,
        languages=languages,
        metadata=document_metadata,
    )


def _read_pdf(path: Path) -> tuple[str, dict]:
    from pypdf import PdfReader

    pages: list[str] = []
    page_offsets: list[int] = []
    cursor = 0
    for page in PdfReader(str(path)).pages:
        page_text = (page.extract_text() or "").strip()
        if not page_text:
            continue
        page_offsets.append(cursor)
        pages.append(page_text)
        cursor += len(page_text) + 2
    return "\n\n".join(pages), {"page_offsets": page_offsets, "page_count": len(pages)}


def _read_docx(path: Path) -> tuple[str, dict]:
    from docx import Document as DocxDocument

    document = DocxDocument(str(path))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
    text = "\n\n".join(paragraph for paragraph in paragraphs if paragraph)
    return text, {"paragraph_count": sum(bool(paragraph) for paragraph in paragraphs)}


def load_document(path: str | Path, *, language: str | None = None) -> Document:
    file_path = Path(path)
    suffix = file_path.suffix.casefold()
    if suffix not in SUPPORTED_SUFFIXES:
        supported = ", ".join(sorted(SUPPORTED_SUFFIXES))
        raise ValueError(f"Unsupported document type {suffix!r}. Supported: {supported}")
    if suffix == ".pdf":
        text, extra = _read_pdf(file_path)
    elif suffix == ".docx":
        text, extra = _read_docx(file_path)
    else:
        text, extra = file_path.read_text(encoding="utf-8"), {}
    metadata = {"suffix": suffix, "size_bytes": file_path.stat().st_size, **extra}
    return document_from_text(file_path.name, text, language=language, metadata=metadata)


def load_directory(path: str | Path) -> list[Document]:
    directory = Path(path)
    return [
        load_document(file_path)
        for file_path in sorted(directory.iterdir())
        if file_path.is_file() and file_path.suffix.casefold() in SUPPORTED_SUFFIXES
    ]


def _page_for_offset(offset: int, page_offsets: list[int]) -> int | None:
    if not page_offsets:
        return None
    return max(1, bisect_right(page_offsets, offset))


def chunk_document(
    document: Document,
    *,
    chunk_size: int = 900,
    overlap: int = 140,
) -> list[Chunk]:
    if chunk_size < 200:
        raise ValueError("chunk_size must be at least 200 characters")
    if overlap < 0 or overlap >= chunk_size:
        raise ValueError("overlap must be non-negative and smaller than chunk_size")

    chunks: list[Chunk] = []
    text = document.text
    position = 0
    page_offsets = list(document.metadata.get("page_offsets", []))

    # Blank-line blocks preserve paragraphs from PDF/DOCX/TXT and Markdown.
    # Short headings are attached to the following paragraph, producing
    # topical passages instead of one file-sized chunk.
    blocks = list(re.finditer(r"\S(?:.*?)(?=\n\s*\n|\Z)", text, re.S))
    passages: list[tuple[int, int]] = []
    pending_start: int | None = None
    for block in blocks:
        block_text = block.group(0).strip()
        is_heading = block_text.startswith("#") or (
            len(block_text) <= 80 and not block_text.endswith((".", "!", "?", "؟", "؛"))
        )
        if is_heading:
            if pending_start is None:
                pending_start = block.start()
            continue
        passages.append((pending_start if pending_start is not None else block.start(), block.end()))
        pending_start = None
    if pending_start is not None:
        passages.append((pending_start, len(text)))
    if not passages and text.strip():
        passages.append((0, len(text)))

    for passage_start, passage_end in passages:
        start = passage_start
        while start < passage_end:
            tentative_end = min(passage_end, start + chunk_size)
            end = tentative_end
            if tentative_end < passage_end:
                paragraph_break = text.rfind("\n", start + chunk_size // 2, tentative_end)
                sentence_break = max(
                    text.rfind(". ", start + chunk_size // 2, tentative_end),
                    text.rfind("؟ ", start + chunk_size // 2, tentative_end),
                    text.rfind("! ", start + chunk_size // 2, tentative_end),
                )
                best_break = max(paragraph_break, sentence_break)
                if best_break > start:
                    end = best_break + 1

            chunk_text = text[start:end].strip()
            if chunk_text:
                languages = detect_languages(chunk_text)
                language_label = (
                    "mixed"
                    if len(languages) > 1
                    else languages[0]
                    if languages
                    else document.language
                )
                metadata = {
                    "start_char": start,
                    "end_char": end,
                    "languages": list(languages or document.languages),
                }
                page = _page_for_offset(start, page_offsets)
                if page is not None:
                    metadata["page"] = page
                chunks.append(
                    Chunk(
                        id=_stable_id(document.id, str(position), chunk_text),
                        document_id=document.id,
                        source=document.name,
                        text=chunk_text,
                        language=language_label,
                        position=position,
                        metadata=metadata,
                    )
                )
                position += 1
            if end >= passage_end:
                break
            start = max(start + 1, end - overlap)
            while start < passage_end and text[start].isspace():
                start += 1
    return chunks


def chunk_documents(documents: Iterable[Document], **kwargs) -> list[Chunk]:
    return [chunk for document in documents for chunk in chunk_document(document, **kwargs)]
